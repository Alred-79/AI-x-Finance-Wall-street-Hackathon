"""Extract the buyer's own scoring rules from the questionnaire workbook → buyer_rules.yaml.

Run: python -m src.app.catalog.build_buyer_rules
Reads: SecurityQuestionnaireMatrix, InherentRiskScore, ResidualRiskScore, Vendor Security Responses,
Vendor Criticality Table, Vendor Profile (requested documents).
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import yaml

from ..config import settings

OUT = settings.catalog_dir / "buyer_rules.yaml"


def _find_workbook() -> Path:
    for p in settings.datasets_dir.rglob("*.xlsx"):
        if "questionnaire" in p.name.lower():
            return p
    raise FileNotFoundError("questionnaire workbook not found under datasets/")


def _norm_qid(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,3})(?:\.0+)?", s)
    return f"{int(m.group(1))}" if m else None


def _header_map(ws, key_header: str) -> tuple[int, dict[str, int]]:
    for r in range(1, min(ws.max_row, 60) + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(isinstance(v, str) and v.strip().startswith(key_header) for v in vals):
            hm = {}
            for c, v in enumerate(vals, start=1):
                if isinstance(v, str) and v.strip():
                    hm[v.strip().split("\n")[0].lower()] = c
            return r, hm
    raise ValueError(f"header '{key_header}' not found in {ws.title}")


def build() -> dict:
    wb = openpyxl.load_workbook(str(_find_workbook()), data_only=True)
    out: dict = {"questions": {}, "criticality_table": [], "requested_documents": [], "attestations": []}

    # --- questions & topics from Vendor Security Responses
    ws = wb["Vendor Security Responses"]
    hr, hm = _header_map(ws, "Security Question ID#")
    topic = ""
    for r in range(hr + 1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and a.strip().lower() == "topic":
            topic = str(b).strip()
            continue
        qid = _norm_qid(a)
        if qid and isinstance(b, str) and b.strip() and b.strip() != str(a):
            out["questions"][qid] = {"qid": qid, "topic": topic, "text": b.strip()}

    # --- matrix: remediation rule, control id, objective, criticality
    ws = wb["SecurityQuestionnaireMatrix"]
    hr, hm = _header_map(ws, "Security Question ID#")
    col = lambda name: next((c for k, c in hm.items() if k.startswith(name)), None)  # noqa: E731
    c_rem, c_cid, c_obj, c_crit = col("remediation"), col("control id"), col("control objective"), col("control criticality")
    c_inh = col("inherent risk score")
    for r in range(hr + 1, ws.max_row + 1):
        qid = _norm_qid(ws.cell(r, 1).value)
        if not qid:
            continue
        q = out["questions"].setdefault(qid, {"qid": qid, "topic": "", "text": str(ws.cell(r, 2).value or "").strip()})
        rule = ws.cell(r, c_rem).value if c_rem else None
        q["rule"] = str(rule).strip() if rule else ""
        q["control_id"] = str(ws.cell(r, c_cid).value or "").strip() if c_cid else ""
        q["control_objective"] = str(ws.cell(r, c_obj).value or "").strip() if c_obj else ""
        q["criticality"] = str(ws.cell(r, c_crit).value or "").strip() if c_crit else ""
        inh = ws.cell(r, c_inh).value if c_inh else None
        q["informational"] = isinstance(inh, str) and "informational" in inh.lower()
        q["escalate_if_no"] = "escalat" in q["rule"].lower()

    # --- inherent / residual scores
    for sheet, key in (("InherentRiskScore", "inherent_pts"), ("ResidualRiskScore", "residual_pts")):
        ws = wb[sheet]
        for r in range(2, ws.max_row + 1):
            qid = _norm_qid(ws.cell(r, 1).value)
            if not qid or qid not in out["questions"]:
                continue
            v = ws.cell(r, 2).value
            note = ws.cell(r, 3).value
            q = out["questions"][qid]
            q[key] = float(v) if isinstance(v, (int, float)) else 0.0
            if isinstance(note, str) and note.strip():
                q.setdefault("score_notes", []).append(note.strip())
                if "informational" in note.lower():
                    q["informational"] = True

    # --- vendor criticality table
    ws = wb["Vendor Criticality Table"]
    for r in range(2, ws.max_row + 1):
        vt, al, sc = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if vt and al and sc:
            out["criticality_table"].append({"vendor_type": str(vt).strip(), "access": str(al).strip(), "criticality": str(sc).strip()})

    # --- requested documents & attestations from Vendor Profile
    ws = wb["Vendor Profile"]
    section = ""
    for r in range(1, ws.max_row + 1):
        a, c = ws.cell(r, 1).value, ws.cell(r, 3).value
        if isinstance(a, str) and "attestations" in a.lower():
            section = "att"
            continue
        if isinstance(a, str) and "documentation" in a.lower():
            section = "doc"
            continue
        if isinstance(a, str) and a.strip() and section:
            section = ""
        if section == "att" and isinstance(c, str) and c.strip():
            out["attestations"].append(c.strip())
        if section == "doc" and isinstance(c, str) and c.strip():
            out["requested_documents"].append(c.strip())
    return out


def main() -> None:
    data = build()
    OUT.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=120))
    qs = data["questions"]
    scored = sum(1 for q in qs.values() if q.get("inherent_pts"))
    print(f"wrote {OUT} — {len(qs)} questions, {scored} scored, {len(data['criticality_table'])} criticality rows, "
          f"{len(data['requested_documents'])} requested docs")


if __name__ == "__main__":
    main()
