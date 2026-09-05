"""File parsers: docx / xlsx / pdf / images / text → ParsedDoc.

Each parser returns plain text with light structure markers so the chunker can split
on headings and keep table rows intact. Images and scanned PDFs are described with a
vision model (lazy; only when an LLM key is available).
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import docx
import pypdf

TEXT_EXT = {".md", ".txt", ".csv", ".log"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}


@dataclass
class ParsedDoc:
    path: str
    name: str
    folder: str
    ext: str
    text: str
    meta: dict = field(default_factory=dict)
    needs_vision: bool = False


def _docx(path: Path) -> str:
    d = docx.Document(str(path))
    parts: list[str] = []
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if "heading" in style or "title" in style:
            parts.append(f"\n## {t}")
        else:
            parts.append(t)
    for ti, table in enumerate(d.tables):
        parts.append(f"\n## Table {ti + 1}")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"\n## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                # trim trailing empties
                while cells and cells[-1] == "":
                    cells.pop()
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf(path: Path) -> tuple[str, bool]:
    r = pypdf.PdfReader(str(path))
    pages = [(pg.extract_text() or "").strip() for pg in r.pages]
    text = "\n\n".join(p for p in pages if p)
    return text, len(text) < 40


def _chat_json(path: Path) -> str:
    """Slack/Teams-style export: list of {user, text, ts, channel}."""
    data = json.loads(path.read_text())
    msgs = data if isinstance(data, list) else data.get("messages", [])
    lines = []
    for m in msgs:
        if not isinstance(m, dict):
            continue
        who = m.get("user") or m.get("user_name") or m.get("author") or "unknown"
        ts = m.get("ts") or m.get("timestamp") or m.get("date") or ""
        ch = m.get("channel") or ""
        text = m.get("text") or m.get("message") or ""
        lines.append(f"[{ts}] {('#' + ch + ' ') if ch else ''}{who}: {text}")
    return "\n".join(lines)


def parse_file(path: Path, root: Path) -> ParsedDoc:
    ext = path.suffix.lower()
    rel = path.relative_to(root) if root in path.parents else path
    folder = rel.parts[0] if len(rel.parts) > 1 else ""
    text, needs_vision = "", False
    if ext == ".docx":
        text = _docx(path)
    elif ext in {".xlsx", ".xlsm"}:
        text = _xlsx(path)
    elif ext == ".pdf":
        text, needs_vision = _pdf(path)
    elif ext == ".json":
        try:
            text = _chat_json(path)
        except Exception:  # noqa: BLE001
            text = path.read_text(errors="ignore")
    elif ext in TEXT_EXT:
        text = path.read_text(errors="ignore")
    elif ext in IMAGE_EXT:
        text, needs_vision = "", True
    else:
        text = ""
    text = re.sub(r"\n{3,}", "\n\n", text)
    return ParsedDoc(path=str(path), name=path.name, folder=folder, ext=ext, text=text, needs_vision=needs_vision)


def iter_files(root: Path) -> list[Path]:
    out = []
    for p in sorted(root.rglob("*")):
        if p.is_file() and not p.name.startswith(("~$", ".")):
            out.append(p)
    return out
