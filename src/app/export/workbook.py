"""Fill the customer's own workbook: Vendor Security Responses, Vendor Profile, Reputational Assessment,
plus two added sheets (Analyst Summary, Exception Requests). Formulas in reviewer tabs are preserved.

Provenance guard: a response is written ONLY if it has evidence (claims / employee statements). Otherwise the cell
gets "Unknown — needs confirmation". This is enforced here, not by the model.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..catalog.loader import get_catalog
from ..config import settings
from ..engine.derive import question_states
from ..engine.scorer import document_availability, score
from ..store.db import Store, loads

FILL = {
    "VERIFIED": "C6EFCE", "CONFIRMED_BY_USER": "DDEBF7", "PARTIAL": "FFEB9C", "CONFLICT": "F8CBAD", "UNKNOWN": "E7E6E6",
}
STATUS_LABEL = {
    "VERIFIED": "Verified from company documents", "CONFIRMED_BY_USER": "Confirmed by employee",
    "PARTIAL": "Partially supported — needs confirmation", "CONFLICT": "Conflicting sources — needs confirmation",
    "UNKNOWN": "Unknown — needs confirmation",
}
REP_ROWS = {
    "vendor_legal_name": 5, "dba_brand_names": 6, "corporate_website": 7, "trust_portal_url": 8, "headquarters": 9,
    "countries_of_operation": 10, "year_founded": 11, "corporate_structure": 13, "bbb_rating": 17,
    "publicly_reported_issues": 18, "news_media_risk_indicators": 19, "social_media_risk_indicators": 20,
    "sanctions_ofac": 21, "sos_standing": 22, "legal_compliance_posture": 23, "assurance_documents_public_claims": 24,
    "regulatory_investigations": 25,
}


def _find_workbook() -> Path:
    for p in settings.datasets_dir.rglob("*.xlsx"):
        if "questionnaire" in p.name.lower():
            return p
    raise FileNotFoundError("questionnaire workbook not found")


def _norm_qid(v) -> str | None:
    if v is None:
        return None
    m = re.fullmatch(r"(\d{1,3})(?:\.0+)?", str(v).strip())
    return m.group(1) if m else None


def _set(ws, row: int, col: int, value, wrap: bool = True) -> None:
    """Write to a cell, redirecting to the anchor cell if (row, col) is inside a merged range."""
    cell = ws.cell(row, col)
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            cell = ws.cell(rng.min_row, rng.min_col)
            break
    cell.value = value
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _evidence_text(ev: list[dict]) -> str:
    parts = []
    for e in ev[:6]:
        if e.get("type") == "user":
            parts.append(f"Confirmed by {e['doc']} on {str(e.get('date', ''))[:10]}: \"{e['statement']}\"")
        elif e.get("type") == "external":
            parts.append(f"Public source: {e.get('excerpt') or e.get('doc')}")
        elif e.get("type") == "template":
            parts.append(f"[template, not evidence] {e['doc']}")
        else:
            parts.append(f"{e['doc']}: \"{(e.get('excerpt') or e.get('statement') or '')[:180]}\"")
    return "\n".join(parts)


def fill_workbook(store: Store, out_path: Path | None = None) -> Path:
    src = _find_workbook()
    wb = openpyxl.load_workbook(str(src))  # keep formulas
    states = {s["qid"]: s for s in question_states(store)}
    cat = get_catalog()

    # --- Vendor Security Responses
    ws = wb["Vendor Security Responses"]
    ws["B1"] = settings.vendor_legal_name
    for r in range(4, ws.max_row + 1):
        qid = _norm_qid(ws.cell(r, 1).value)
        if not qid or qid not in states:
            continue
        s = states[qid]
        has_evidence = bool(s["evidence"]) and any(e.get("type") in ("claim", "user") for e in s["evidence"])
        if s["status"] == "UNKNOWN" or not has_evidence:
            resp, comments = "Unknown — needs confirmation", (s["comments"] or "No supporting evidence found in company data.")
        else:
            resp, comments = s["answer"], s["comments"]
        prefix = STATUS_LABEL[s["status"]] + (f" (confidence {int(round(s['confidence'] * 100))}%)" if s["confidence"] else "")
        ws.cell(r, 3).value = resp
        ws.cell(r, 4).value = f"[{prefix}] {comments}".strip()
        ws.cell(r, 5).value = _evidence_text(s["evidence"]) or "—"
        for c in (3, 4, 5):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=FILL[s["status"]])
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in (("C", 60), ("D", 70), ("E", 70)):
        ws.column_dimensions[col].width = width

    # --- Vendor Profile: attestations, documents, AI
    vp = wb["Vendor Profile"]
    vp["C5"] = settings.vendor_legal_name
    docs = {d["document"]: d for d in document_availability(store)}
    rep = store.kv_get("reputational") or {}
    att_claims = store.q("SELECT statement FROM claims WHERE control='company_profile' AND (statement LIKE ? OR statement LIKE ?)", ("%SOC 2%", "%ISO 27001%"))
    soc2_doc = store.one("SELECT name FROM documents WHERE doc_type='soc2_report'")
    for r in range(10, 16):
        label = str(vp.cell(r, 3).value or "")
        if label == "SOC 2 Type 2" and soc2_doc:
            vp.cell(r, 2).value = True
            vp.cell(r, 5).value = f"Report on file: {soc2_doc['name']}. NOTE: verify entity/period/CEO named in the report before attaching (analyst flagged a discrepancy)." if any("Prasun" in c["statement"] or "CEO" in c["statement"] for c in att_claims) else f"Report on file: {soc2_doc['name']}"
        elif label == "ISO 27001":
            vp.cell(r, 2).value = False
            vp.cell(r, 5).value = "No ISO 27001 certificate found in company data."
    for r in range(17, 25):
        label = str(vp.cell(r, 3).value or "")
        d = docs.get(label)
        if d:
            vp.cell(r, 2).value = d["status"] == "available"
            vp.cell(r, 5).value = {"available": f"Available: {d['note']}", "partial": f"Justification: {d['note']}",
                                   "template_only": f"Justification: {d['note']}", "missing": f"Justification required: {d['note']}"}[d["status"]]
    vp["B29"] = True
    vp["D29"] = "Yes — the product is an AI-native compliance automation platform (AI agents collect and validate evidence)."
    vp["B31"] = True
    vp["D31"] = "Client data is processed by AI features; the vendor must attest to no model training on client data (buyer checklist item 5) — confirm with the CTO."

    # --- Reputational Assessment (from outside-in research)
    ra = wb["Reputational Assessment"]
    fields = (rep.get("fields") or {})
    sources = rep.get("sources") or {}
    for key, row in REP_ROWS.items():
        val = fields.get(key) or ("Not researched yet" if not rep else "Not found in public sources")
        src_ids = sources.get(key) or []
        _set(ra, row, 2, val + (f"  [sources: {', '.join(src_ids)}]" if src_ids else ""))
    if fields.get("year_founded") and re.search(r"\d{4}", str(fields["year_founded"])):
        _set(ra, 12, 2, date.today().year - int(re.search(r"\d{4}", str(fields["year_founded"])).group()))
    ra.column_dimensions["B"].width = 90

    # --- Analyst Summary sheet
    sm = wb.create_sheet("Analyst Summary")
    sc = store.kv_get("last_score") or score(store)
    head = ["Q#", "Topic", "Question", "Status", "Answer value", "Confidence", "Buyer criticality", "Inherent pts", "Predicted escalation", "Open slots / next question"]
    sm.append(head)
    for c in range(1, len(head) + 1):
        sm.cell(1, c).font = Font(bold=True)
    esc = {e["qid"] for e in sc["escalations"] if e.get("qid")}
    for q in cat.ordered():
        s = states[q.qid]
        sm.append([q.qid, q.topic, q.text, STATUS_LABEL[s["status"]], s["answer_value"], s["confidence"], q.criticality, q.inherent_pts,
                   "Yes" if q.qid in esc else "", "; ".join(s["open_slots"].keys()) + ((" → " + s["next_question"]) if s["next_question"] else "")])
        sm.cell(sm.max_row, 4).fill = PatternFill("solid", fgColor=FILL[s["status"]])
    sm.append([])
    sm.append(["Buyer's-eye summary"])
    sm.cell(sm.max_row, 1).font = Font(bold=True)
    sm.append(["Vendor criticality", sc["vendor_criticality"]])
    sm.append(["Predicted inherent risk points", f"{sc['inherent_points']} / {sc['max_inherent_points']}"])
    sm.append(["Predicted rating", sc["predicted_rating"]])
    sm.append(["Predicted escalations", len(sc["escalations"])])
    for e in sc["escalations"]:
        sm.append(["", f"Q{e['qid']}" if e.get("qid") else "Document", e["text"], e["reason"]])
    for i, w in enumerate((6, 26, 60, 34, 14, 11, 16, 12, 12, 70), start=1):
        sm.column_dimensions[get_column_letter(i)].width = w

    # --- Exception Requests sheet
    ex = wb.create_sheet("Exception Requests")
    ex.append(["Q#", "Question", "Vendor response", "Justification", "Mitigating controls", "Remediation plan", "Evidence to attach", "Control ID", "Criticality"])
    for c in range(1, 10):
        ex.cell(1, c).font = Font(bold=True)
    for e in store.kv_get("exceptions") or []:
        q = cat.questions.get(e["qid"])
        s = states.get(e["qid"], {})
        ex.append([e["qid"], e["text"], s.get("answer", ""), e.get("justification", ""), e.get("mitigating_controls", ""), e.get("remediation_plan", ""),
                   e.get("evidence_to_attach", ""), q.control_id if q else "", q.criticality if q else ""])
    for i, w in enumerate((6, 50, 40, 60, 50, 50, 40, 12, 12), start=1):
        ex.column_dimensions[get_column_letter(i)].width = w
    for row in ex.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = out_path or (settings.data_dir / "exports" / f"{settings.vendor_brand}_Security_Questionnaire_completed_{date.today().isoformat()}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    store.log("export_workbook", str(out))
    return out
