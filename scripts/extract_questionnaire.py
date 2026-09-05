"""Extract the organizer's vendor security questionnaire into structured JSON.

Two sheets of the real Regodit workbook are merged:

  'Vendor Security Responses'    -> question id, topic grouping, question text
  'SecurityQuestionnaireMatrix'  -> Control ID, control objective, criticality,
                                    and the reviewer's own escalation playbook
                                    ("If answer is No: ... escalate")

Merging them means every answer the analyst produces is tied to the sponsor's
control framework and to the exact reviewer action that follows from it.
"""
from __future__ import annotations

import json
import re
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parents[1]
SRC = (
    ROOT
    / "data"
    / "raw"
    / "1. Sample_Vendor questionnaire"
    / "Regodit_Comprehensive_Vendor_Security_Questionnaire_Clean.xlsx"
)
OUT = ROOT / "data" / "questionnaire.json"

RESPONSES_SHEET = "Vendor Security Responses"
MATRIX_SHEET = "SecurityQuestionnaireMatrix"

# Column indexes in SecurityQuestionnaireMatrix (header row 18 of the real file).
M_ID, M_QUESTION, M_PLAYBOOK = 0, 1, 4
M_CONTROL_ID, M_CONTROL_OBJ, M_CRITICALITY = 14, 15, 16

BAD_VALUES = {"", "#REF!", "#NAME?", "#N/A", "N/A", "Not Escalated"}

# Questions whose answer is a document/attachment request rather than a control claim.
ATTACHMENT_RE = re.compile(r"please attach|request (a )?copy|provide (a )?copy", re.I)


def clean(value: object) -> str:
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return "" if text in BAD_VALUES else text


def norm_id(raw: str) -> str:
    """'1' / '1.0' -> '1.0' so the two sheets join reliably."""
    try:
        return f"{float(raw):.1f}"
    except ValueError:
        return raw


def read_matrix(wb: openpyxl.Workbook) -> dict[str, dict]:
    ws = wb[MATRIX_SHEET]
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=19, values_only=True):
        if row is None or len(row) <= M_CONTROL_OBJ:
            continue
        qid = clean(row[M_ID])
        if not re.fullmatch(r"\d+(\.\d+)?", qid):
            continue
        criticality = clean(row[M_CRITICALITY]) if len(row) > M_CRITICALITY else ""
        out[norm_id(qid)] = {
            "control_id": clean(row[M_CONTROL_ID]),
            "control_objective": clean(row[M_CONTROL_OBJ]),
            "criticality": criticality,
            "reviewer_playbook": clean(row[M_PLAYBOOK]),
        }
    return out


def read_questions(wb: openpyxl.Workbook) -> list[dict]:
    ws = wb[RESPONSES_SHEET]
    questions: list[dict] = []
    topic = "General"
    for row in ws.iter_rows(min_row=4, values_only=True):
        col_a, col_b = clean(row[0]), clean(row[1])
        if not col_a and not col_b:
            continue
        if col_a.lower() == "topic":
            if col_b:
                topic = col_b
            continue
        if re.fullmatch(r"\d+(\.\d+)?", col_a) and col_b:
            questions.append({"id": norm_id(col_a), "topic": topic, "question": col_b})
    return questions


def priority_of(q: dict) -> int:
    """Investigation priority: 1 = highest. Drives 'ask the important things first'."""
    crit = (q.get("criticality") or "").lower()
    if crit.startswith("high") or crit.startswith("critical"):
        return 1
    if crit.startswith("med"):
        return 2
    if q.get("control_id"):
        return 3
    return 4


def main() -> None:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    matrix = read_matrix(wb)
    questions = read_questions(wb)

    matched = 0
    for q in questions:
        meta = matrix.get(q["id"])
        if meta:
            matched += 1
            q.update(meta)
        else:
            q.update(
                {
                    "control_id": "",
                    "control_objective": "",
                    "criticality": "",
                    "reviewer_playbook": "",
                }
            )
        q["answer_kind"] = "attachment" if ATTACHMENT_RE.search(q["question"]) else "control"
        q["priority"] = priority_of(q)

    payload = {
        "source_file": SRC.name,
        "source_sheets": [RESPONSES_SHEET, MATRIX_SHEET],
        "question_count": len(questions),
        "matrix_matched": matched,
        "topics": sorted({q["topic"] for q in questions}),
        "questions": questions,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    with_control = sum(1 for q in questions if q["control_id"])
    print(f"Extracted {len(questions)} questions across {len(payload['topics'])} topics")
    print(f"  matrix rows joined : {matched}")
    print(f"  with Control ID    : {with_control}")
    print(f"  attachment requests: {sum(1 for q in questions if q['answer_kind'] == 'attachment')}")
    print("  priority spread    : " + ", ".join(
        f"P{p}={sum(1 for q in questions if q['priority'] == p)}" for p in (1, 2, 3, 4)
    ))
    print(f"\nWrote {OUT}\n")
    for q in questions[:5]:
        print(f"  [{q['id']}] P{q['priority']} {q['control_id'] or '---':10s} {q['question'][:80]}")


if __name__ == "__main__":
    main()
