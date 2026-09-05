"""One-off explorer: dump structure of the organizer-provided dataset."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from docx import Document

RAW = Path(__file__).resolve().parents[1] / "data" / "raw"


def peek_xlsx(path: Path, max_rows: int = 12) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"\n=== XLSX {path.name} sheets={wb.sheetnames}")
    for ws in wb.worksheets:
        print(f"--- sheet '{ws.title}' dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                print("    ...")
                break
            cells = ["" if c is None else str(c).replace("\n", " ")[:70] for c in row]
            if any(cells):
                print(f"    {i}: {cells}")


def peek_docx(path: Path, max_paras: int = 30) -> None:
    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    print(f"\n=== DOCX {path.name} paras={len(paras)} tables={len(doc.tables)}")
    for p in paras[:max_paras]:
        print(f"    {p[:150]}")
    for ti, table in enumerate(doc.tables[:2]):
        print(f"    -- table {ti} rows={len(table.rows)}")
        for row in table.rows[:6]:
            print(f"       {[c.text.strip()[:45] for c in row.cells]}")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "all"
    for path in sorted(RAW.rglob("*")):
        if path.is_dir():
            continue
        if target != "all" and target.lower() not in path.name.lower():
            continue
        if path.suffix.lower() == ".xlsx":
            peek_xlsx(path)
        elif path.suffix.lower() == ".docx":
            peek_docx(path)
