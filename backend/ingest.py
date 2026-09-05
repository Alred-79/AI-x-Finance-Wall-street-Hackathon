"""Ingest the organizer-provided company corpus into citable evidence chunks.

Every chunk keeps a precise, human-checkable citation:
  source_file, source_category, locator (heading / sheet!row / page)

Supported: .docx (paragraphs + tables), .xlsx (row records), .pdf (page text),
images (registered as non-text evidence that a human must review — we never
pretend to have read a diagram).
"""
from __future__ import annotations

import json
import re
import warnings
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from .config import CORPUS_PATH, RAW_DIR

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

MAX_CHARS = 1400
MIN_CHARS = 40

# Category names come from the organizer's folder numbering.
CATEGORY_MAP = {
    "1. Sample_Vendor questionnaire": "questionnaire",
    "2. Company policies": "policy",
    "3. Security Assessment Reports": "assessment_report",
    "4. Contracts_agreements": "contract",
    "5. Infrastructure_internal info": "infrastructure",
}

# Sheets that are template scaffolding, not evidence.
SKIP_SHEET_TOKENS = ("quicklink", "definitions", "instructions")


@dataclass
class Chunk:
    id: str
    text: str
    source_file: str
    source_category: str
    locator: str
    doc_title: str
    kind: str = "text"
    tags: list[str] = field(default_factory=list)


def _norm(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text).strip()


def _category_of(path: Path) -> str:
    for part in path.parts:
        if part in CATEGORY_MAP:
            return CATEGORY_MAP[part]
    return "other"


def _doc_title(path: Path) -> str:
    return re.sub(r"[_\-]+", " ", path.stem).strip()


def _is_heading(par: Paragraph) -> bool:
    style = (par.style.name or "").lower()
    if style.startswith("heading") or style in {"title", "subtitle"}:
        return True
    text = par.text.strip()
    if not text or len(text) > 90:
        return False
    # Numbered clause headings such as "4.2 Access Control" or ALL-CAPS titles.
    if re.match(r"^\d+(\.\d+)*\s+\S", text) and text == text.rstrip("."):
        return True
    return text.isupper() and len(text.split()) <= 12


def _iter_docx_blocks(doc: Document):
    """Yield paragraphs and tables in true document order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def _flush(buf: list[str], heading: str, path: Path, out: list[Chunk]) -> None:
    text = _norm("\n".join(buf))
    if len(text) < MIN_CHARS:
        return
    idx = len(out)
    out.append(
        Chunk(
            id=f"{path.stem}::{idx}",
            text=text,
            source_file=str(path.relative_to(RAW_DIR)).replace("\\", "/"),
            source_category=_category_of(path),
            locator=heading or "body",
            doc_title=_doc_title(path),
        )
    )


def parse_docx(path: Path) -> list[Chunk]:
    doc = Document(str(path))
    out: list[Chunk] = []
    heading = ""
    buf: list[str] = []

    for block in _iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            text = _norm(block.text)
            if not text:
                continue
            if _is_heading(block):
                _flush(buf, heading, path, out)
                buf = []
                heading = text
                continue
            buf.append(text)
            if sum(len(b) for b in buf) > MAX_CHARS:
                _flush(buf, heading, path, out)
                buf = []
        else:  # Table
            rows: list[str] = []
            for row in block.rows:
                cells = [_norm(c.text) for c in row.cells]
                # collapse repeated merged-cell text
                dedup: list[str] = []
                for c in cells:
                    if not dedup or dedup[-1] != c:
                        dedup.append(c)
                line = " | ".join(c for c in dedup if c)
                if line:
                    rows.append(line)
            if rows:
                _flush(buf, heading, path, out)
                buf = []
                block_text = "\n".join(rows)
                for i in range(0, len(block_text), MAX_CHARS):
                    piece = block_text[i : i + MAX_CHARS]
                    idx = len(out)
                    out.append(
                        Chunk(
                            id=f"{path.stem}::t{idx}",
                            text=piece,
                            source_file=str(path.relative_to(RAW_DIR)).replace("\\", "/"),
                            source_category=_category_of(path),
                            locator=f"{heading or 'table'} (table)",
                            doc_title=_doc_title(path),
                            kind="table",
                        )
                    )
    _flush(buf, heading, path, out)
    return out


def parse_xlsx(path: Path) -> list[Chunk]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[Chunk] = []
    rel = str(path.relative_to(RAW_DIR)).replace("\\", "/")

    for ws in wb.worksheets:
        title_l = ws.title.lower()
        if any(tok in title_l for tok in SKIP_SHEET_TOKENS):
            continue

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Find the header row: first row with >=2 non-empty string cells.
        header: list[str] = []
        header_at = -1
        for i, row in enumerate(rows[:12]):
            cells = [_norm(str(c)) if c is not None else "" for c in row]
            filled = [c for c in cells if c]
            if len(filled) >= 2 and sum(len(c) for c in filled) > 12:
                header = cells
                header_at = i
                break
        if header_at < 0:
            header_at = 0

        for r, row in enumerate(rows[header_at + 1 :], start=header_at + 2):
            cells = [_norm(str(c)) if c is not None else "" for c in row]
            if not any(cells):
                continue
            pairs = []
            for h, c in zip(header, cells):
                if not c or c in {"N/A", "Not Escalated", "#REF!"}:
                    continue
                pairs.append(f"{h}: {c}" if h else c)
            if not pairs:
                continue
            text = _norm("; ".join(pairs))
            if len(text) < MIN_CHARS:
                continue
            out.append(
                Chunk(
                    id=f"{path.stem}::{ws.title}::r{r}",
                    text=text,
                    source_file=rel,
                    source_category=_category_of(path),
                    locator=f"{ws.title}!row {r}",
                    doc_title=_doc_title(path),
                    kind="record",
                )
            )
    return out


def parse_pdf(path: Path) -> list[Chunk]:
    out: list[Chunk] = []
    rel = str(path.relative_to(RAW_DIR)).replace("\\", "/")
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pragma: no cover
        print(f"  ! cannot read {path.name}: {exc}")
        return out

    for pno, page in enumerate(reader.pages, start=1):
        try:
            text = _norm(page.extract_text() or "")
        except Exception:
            text = ""
        if len(text) < MIN_CHARS:
            continue
        for i in range(0, len(text), MAX_CHARS):
            piece = text[i : i + MAX_CHARS]
            out.append(
                Chunk(
                    id=f"{path.stem}::p{pno}::{i}",
                    text=piece,
                    source_file=rel,
                    source_category=_category_of(path),
                    locator=f"page {pno}",
                    doc_title=_doc_title(path),
                )
            )
    return out


def register_image(path: Path) -> list[Chunk]:
    """Images are registered, never interpreted. Marked for human review."""
    rel = str(path.relative_to(RAW_DIR)).replace("\\", "/")
    return [
        Chunk(
            id=f"{path.stem}::image",
            text=(
                f"Non-text evidence artifact on file: '{path.name}'. This is a diagram image "
                f"({_doc_title(path)}). Its contents have NOT been machine-read; a human "
                "reviewer must confirm what it shows before it is cited as proof of a control."
            ),
            source_file=rel,
            source_category=_category_of(path),
            locator="image file",
            doc_title=_doc_title(path),
            kind="image_unread",
            tags=["needs_human_review"],
        )
    ]


def build_corpus() -> dict:
    chunks: list[Chunk] = []
    files_seen: list[dict] = []

    for path in sorted(RAW_DIR.rglob("*")):
        if path.is_dir() or path.name.startswith("~$"):
            continue
        suffix = path.suffix.lower()
        if suffix == ".docx":
            new = parse_docx(path)
        elif suffix == ".xlsx":
            new = parse_xlsx(path)
        elif suffix == ".pdf":
            new = parse_pdf(path)
        elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
            new = register_image(path)
        else:
            continue
        chunks.extend(new)
        files_seen.append(
            {
                "file": str(path.relative_to(RAW_DIR)).replace("\\", "/"),
                "category": _category_of(path),
                "chunks": len(new),
            }
        )
        print(f"  {len(new):5d} chunks  {path.relative_to(RAW_DIR)}")

    payload = {
        "chunk_count": len(chunks),
        "file_count": len(files_seen),
        "files": files_seen,
        "chunks": [asdict(c) for c in chunks],
    }
    CORPUS_PATH.parent.mkdir(parents=True, exist_ok=True)
    CORPUS_PATH.write_text(json.dumps(payload, indent=1), encoding="utf-8")
    return payload


def main() -> None:
    print(f"Ingesting corpus from {RAW_DIR}")
    payload = build_corpus()
    print(
        f"\nIndexed {payload['chunk_count']} evidence chunks "
        f"from {payload['file_count']} files -> {CORPUS_PATH}"
    )


if __name__ == "__main__":
    main()
